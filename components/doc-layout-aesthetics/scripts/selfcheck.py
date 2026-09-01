# -*- coding: utf-8 -*-
"""冒烟自检：验证依赖、字体、引擎探测、表格跨页保护逻辑与交付物产物。

用途：构建/发布前快速确认环境与核心逻辑健康，逐项打印 ✅/❌ 并返回退出码。

用法：
  python scripts/selfcheck.py [--out DIR] [--build]

  --out DIR   交付物所在目录（默认当前目录），用于产物存在性检查。
  --build     额外真实跑一遍 build_all.py（含 PDF）到临时目录，做端到端构建冒烟。

退出码：0 全部通过；1 存在失败项。

说明：本脚本本身只依赖标准库 + 技能运行时的第三方依赖（docx/reportlab 等），
由 uv 虚拟环境提供（`uv run --project <技能根目录> python scripts/selfcheck.py`）。
"""
import argparse
import os
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)  # 让 import build_docx / build_pdf 可解析

DOCX_NAME = "版面美学观_侯捷Word排版艺术提炼.docx"
PPTX_NAME = "版面美学观_对照演示.pptx"
HTML_NAME = "版面美学观_响应式参考.html"
PDF_NAME = "版面美学观_侯捷Word排版艺术提炼.pdf"


def _ensure_stdout_utf8():
    # 刻意保留在本文件（不共享 _common.ensure_stdout_utf8）：selfcheck 必须在
    # 第三方依赖缺失时仍能启动并报告「依赖缺失」，故保持纯标准库可导入。
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass


class Checker:
    def __init__(self):
        self.fail = 0
        self.passed = 0

    def ok(self, msg):
        self.passed += 1
        print("✅", msg)

    def bad(self, msg):
        self.fail += 1
        print("❌", msg)


def check_deps(c):
    """关键第三方依赖可导入。"""
    mods = ["docx", "pptx", "reportlab"]
    missing = []
    for m in mods:
        try:
            __import__(m)
        except ImportError as e:
            missing.append("%s(%s)" % (m, e))
    if missing:
        c.bad("依赖缺失: " + ", ".join(missing))
    else:
        c.ok("依赖可导入: " + ", ".join(mods))


def check_scripts_importable(c):
    """build_docx / build_pdf 可被 import 且无构建副作用（重构后的关键保障）。"""
    try:
        import build_docx as bd  # noqa: F401
        import build_pdf as bp  # noqa: F401
    except Exception as e:
        c.bad("脚本 import 失败: %s" % e)
        return None, None
    c.ok("build_docx / build_pdf 可安全 import")
    return bd, bp


def check_fonts(c, bp):
    """中文字体可探测（纯 Python PDF 兜底的前提）。"""
    try:
        f = bp.detect_fonts()
    except Exception as e:
        c.bad("字体探测抛异常: %s" % e)
        return
    if f:
        c.ok("中文字体可探测: 正文=%s 标题=%s" % (f[0], f[2]))
    else:
        c.bad("未探测到中文字体（macOS/Linux/Windows 应自带）")


def check_engine(c, bp):
    """引擎探测返回合法值，且回退链自洽。"""
    valid = {"libreoffice", "docx2pdf", "wps", "reportlab"}
    try:
        auto = bp.detect_engine()[0]
    except Exception as e:
        c.bad("detect_engine 抛异常: %s" % e)
        return
    if auto in valid:
        c.ok("detect_engine(auto) -> %s" % auto)
    else:
        c.bad("detect_engine 返回非法值: %r" % auto)
    # resolve_engine 各显式引擎均返回合法值
    for eng in valid:
        try:
            got = bp.resolve_engine(eng)
        except Exception as e:
            c.bad("resolve_engine(%s) 抛异常: %s" % (eng, e))
            continue
        if got in valid:
            pass
        else:
            c.bad("resolve_engine(%s) 返回非法值: %r" % (eng, got))
    c.ok("resolve_engine 回退链自洽（四引擎均返回合法值）")


def check_table_protection(c, bd):
    """protect_table 能正确写入 tblHeader / cantSplit / keepNext 等 XML 标记。"""
    from docx import Document
    from docx.oxml.ns import qn

    try:
        # long 模式：表头 + 全部行禁断
        d = Document()
        t = d.add_table(rows=5, cols=3)
        for row in t.rows:
            for cell in row.cells:
                cell.paragraphs[0].text = "x"
        bd.protect_table(t, mode="long")
        hdr = t.rows[0]._tr.get_or_add_trPr().find(qn("w:tblHeader"))
        assert hdr is not None and hdr.get(qn("w:val")) == "true", "表头 tblHeader 缺失"
        for i, row in enumerate(t.rows):
            cs = row._tr.get_or_add_trPr().find(qn("w:cantSplit"))
            assert cs is not None and cs.get(qn("w:val")) == "true", "第 %d 行 cantSplit 缺失" % i

        # short 模式：整表 keepNext + keepLines
        d2 = Document()
        prev = d2.add_paragraph("标题")
        t2 = d2.add_table(rows=2, cols=2)
        for row in t2.rows:
            for cell in row.cells:
                cell.paragraphs[0].text = "y"
        bd.protect_table(t2, mode="short")
        for row in t2.rows:
            for cell in row.cells:
                pPr = cell.paragraphs[0]._p.get_or_add_pPr()
                assert pPr.find(qn("w:keepNext")) is not None, "keepNext 缺失"
                assert pPr.find(qn("w:keepLines")) is not None, "keepLines 缺失"
        assert prev._p.get_or_add_pPr().find(qn("w:keepNext")) is not None, "前段 keepNext 缺失"
    except AssertionError as e:
        c.bad("表格跨页保护失败: %s" % e)
        return
    except Exception as e:
        c.bad("表格跨页保护抛异常: %s" % e)
        return
    c.ok("protect_table long/short 两种模式写入正确 XML 标记")


def check_artifacts(c, out_dir):
    """交付物产物存在且非空。"""
    for name in (DOCX_NAME, PPTX_NAME, HTML_NAME, PDF_NAME):
        p = os.path.join(out_dir, name)
        if os.path.exists(p) and os.path.getsize(p) > 0:
            c.ok("产物存在: %s (%d bytes)" % (name, os.path.getsize(p)))
        else:
            c.bad("产物缺失或为空: %s" % name)


def check_build(c):
    """端到端构建冒烟：真实跑 build_all.py（含 PDF）到临时目录。"""
    try:
        tmp = tempfile.mkdtemp(prefix="doclayout_selfcheck_")
        cmd = [sys.executable, os.path.join(HERE, "build_all.py"), "--out", tmp, "--pdf"]
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
        if r.returncode != 0:
            c.bad("build_all 失败(退出码 %s)\n%s" % (r.returncode, (r.stdout + r.stderr)[-800:]))
            return
        for name in (DOCX_NAME, PPTX_NAME, HTML_NAME, PDF_NAME):
            p = os.path.join(tmp, name)
            if not (os.path.exists(p) and os.path.getsize(p) > 0):
                c.bad("构建产物缺失: %s" % name)
                return
        c.ok("端到端构建成功：docx/pptx/html/pdf 四产物齐全")
    except Exception as e:
        c.bad("端到端构建抛异常: %s" % e)


def main():
    _ensure_stdout_utf8()
    ap = argparse.ArgumentParser(description="冒烟自检")
    ap.add_argument("--out", default=None, help="交付物目录（默认当前目录）")
    ap.add_argument("--build", action="store_true", help="额外端到端构建冒烟")
    args = ap.parse_args()

    c = Checker()
    print("== doc-layout-aesthetics 冒烟自检 ==")
    check_deps(c)
    bd, bp = check_scripts_importable(c)
    if bp is not None:
        check_fonts(c, bp)
        check_engine(c, bp)
    if bd is not None:
        check_table_protection(c, bd)
    # 产物存在性检查：仅当显式指定 --out，或未启用 --build 时才检查（默认 cwd）；
    # 否则 --build 的端到端构建已在 check_build 里自检产物齐全，避免对 cwd 误报。
    out_dir = args.out or os.getcwd()
    if args.out or not args.build:
        check_artifacts(c, out_dir)
    if args.build:
        check_build(c)

    print("---")
    print("通过 %d 项，失败 %d 项" % (c.passed, c.fail))
    if c.fail:
        print("结果：FAIL")
        sys.exit(1)
    print("结果：PASS")
    sys.exit(0)


if __name__ == "__main__":
    main()
