# -*- coding: utf-8 -*-
"""测试 protect_table：docx 表格跨页保护的两种模式（long/short）。"""
from docx import Document
from docx.oxml.ns import qn

from build_docx import protect_table


def _make_table(rows, cols):
    doc = Document()
    t = doc.add_table(rows=rows, cols=cols)
    for row in t.rows:
        for cell in row.cells:
            cell.paragraphs[0].text = "x"
    return t


def _trPr(row):
    return row._tr.get_or_add_trPr()


def test_long_table_sets_tblheader_only_on_first_row():
    t = _make_table(5, 3)
    protect_table(t, mode="long")
    hdr = _trPr(t.rows[0]).find(qn("w:tblHeader"))
    assert hdr is not None, "表头行缺 tblHeader"
    assert hdr.get(qn("w:val")) == "true"
    for i, row in enumerate(t.rows[1:], start=1):
        assert _trPr(row).find(qn("w:tblHeader")) is None, f"第 {i} 行不应有 tblHeader"


def test_long_table_sets_cantsplit_on_all_rows():
    t = _make_table(5, 3)
    protect_table(t, mode="long")
    for i, row in enumerate(t.rows):
        cs = _trPr(row).find(qn("w:cantSplit"))
        assert cs is not None, f"第 {i} 行缺 cantSplit"
        assert cs.get(qn("w:val")) == "true"


def test_short_table_keeps_cells_together():
    t = _make_table(2, 2)
    protect_table(t, mode="short")
    for row in t.rows:
        for cell in row.cells:
            pPr = cell.paragraphs[0]._p.get_or_add_pPr()
            assert pPr.find(qn("w:keepNext")) is not None, "keepNext 缺失"
            assert pPr.find(qn("w:keepLines")) is not None, "keepLines 缺失"


def test_short_table_keeps_prev_paragraph_together():
    doc = Document()
    p = doc.add_paragraph("标题")
    t = doc.add_table(rows=2, cols=2)
    for row in t.rows:
        for cell in row.cells:
            cell.paragraphs[0].text = "x"
    protect_table(t, mode="short")
    pPr = p._p.get_or_add_pPr()
    assert pPr.find(qn("w:keepNext")) is not None, "前一段缺 keepNext"


def test_short_table_does_not_set_cantsplit():
    t = _make_table(2, 2)
    protect_table(t, mode="short")
    for row in t.rows:
        assert _trPr(row).find(qn("w:cantSplit")) is None


def test_protect_table_idempotent():
    t = _make_table(3, 2)
    protect_table(t, mode="long")
    protect_table(t, mode="long")  # 重复调用不应产生重复元素或异常
    assert len(_trPr(t.rows[0]).findall(qn("w:tblHeader"))) == 1
    for row in t.rows:
        assert len(_trPr(row).findall(qn("w:cantSplit"))) == 1
